from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import pandas as pd
import io
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from cv_analyzer import analyze_cv_bytes, clean_text, load_models
from job_api import fetch_live_jobs
from matcher import calculate_match_score

app = FastAPI(title="JobMatch AI API", version="5.0.0")

# السماح للفرونت إند (Next.js / React) بالاتصال بالباك إند بدون مشاكل CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Pydantic Models ---
class ManualMatchRequest(BaseModel):
    skills: str
    experience_years: float = 2.0
    career_level: str = "Entry Level"
    country_code: str = "eg"

class CVMatchRequest(BaseModel):
    skills: list[str]
    skills_text: str
    cleaned_text: str = ""
    role: str
    experience_level: float
    career_level: str
    country_code: str = "eg"

class CVBuilderRequest(BaseModel):
    name: str
    title: str
    email: str
    phone: str
    linkedin: str
    summary: str
    experience: str
    projects: str
    education: str
    skills: str
    languages: str

# --- Endpoints ---

@app.get("/")
def health_check():
    return {"status": "online", "version": "5.0.0", "message": "JobMatch AI FastAPI Backend is running!"}

@app.post("/api/analyze-cv")
async def analyze_cv_endpoint(file: UploadFile = File(...)):
    file_bytes = await file.read()
    result = analyze_cv_bytes(file_bytes, file.filename)
    if result is None:
        raise HTTPException(status_code=500, detail="Failed to load ML models. Check artifacts_v5 folder.")
    return result

@app.post("/api/match-cv-jobs")
def match_cv_jobs(req: CVMatchRequest):
    live_jobs = fetch_live_jobs(req.role, req.country_code)
    if not live_jobs:
        return {"jobs": []}
    jobs_df = pd.DataFrame(live_jobs)
    ranked_jobs = calculate_match_score(req.model_dump(), jobs_df)
    return {"jobs": ranked_jobs[:12]}

@app.post("/api/manual-recommendations")
def manual_recommendations(req: ManualMatchRequest):
    cleaned_manual = clean_text(req.skills)
    _, _, skills_vectorizer, _ = load_models()
    extracted_list = [s.strip().lower() for s in req.skills.split(",") if s.strip()]
    
    if skills_vectorizer:
        try:
            vec = skills_vectorizer.transform([cleaned_manual])
            vec_skills = skills_vectorizer.inverse_transform(vec)[0].tolist()
            if vec_skills:
                extracted_list = vec_skills
        except Exception:
            pass
            
    cv_data = {
        "skills_text": " ; ".join(extracted_list) if extracted_list else cleaned_manual,
        "skills": extracted_list,
        "experience_level": req.experience_years,
        "career_level": req.career_level,
        "role": "Tech Professional"
    }
    
    live_jobs = fetch_live_jobs(req.skills, req.country_code)
    if not live_jobs:
        return {"jobs": []}
    
    jobs_df = pd.DataFrame(live_jobs)
    ranked_jobs = calculate_match_score(cv_data, jobs_df)
    return {"jobs": ranked_jobs[:12]}

@app.get("/api/live-jobs")
def get_live_jobs(query: str = "Software Engineer", country: str = "eg"):
    jobs = fetch_live_jobs(query, country)
    return {"jobs": jobs[:15]}

@app.post("/api/build-cv")
def build_cv_docx(req: CVBuilderRequest):
    doc = docx.Document()
    
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(req.name or "Your Name")
    name_run.bold = True
    name_run.font.size = Pt(22)
    
    contact_para = doc.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_para.add_run(f"{req.email} | {req.phone} | {req.linkedin}").font.size = Pt(10)
    
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.add_run(req.title).font.size = Pt(12)

    def add_section(heading, text_content):
        if text_content and text_content.strip():
            doc.add_heading(heading, level=2)
            for line in text_content.split('\n'):
                if line.strip():
                    style = 'List Bullet' if heading in ['EXPERIENCE', 'PROJECTS', 'EDUCATION'] else None
                    doc.add_paragraph(line.strip(), style=style)

    add_section('PROFESSIONAL SUMMARY', req.summary)
    add_section('EXPERIENCE', req.experience)
    add_section('PROJECTS', req.projects)
    add_section('EDUCATION', req.education)
    add_section('TECHNICAL SKILLS', req.skills)
    add_section('LANGUAGES', req.languages)
    
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    
    filename = f"{req.name.replace(' ', '_')}_CV.docx" if req.name else "Resume.docx"
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )