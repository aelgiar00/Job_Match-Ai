import streamlit as st
import pandas as pd
import re 
import io
import datetime
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from job_api import fetch_live_jobs
from cv_analyzer import load_models, clean_text, analyze_cv
from matcher import calculate_match_score

st.set_page_config(page_title="JobMatch AI | Career Platform", page_icon="🎯", layout="wide", initial_sidebar_state="expanded")

# --- Session State ---
if 'cv_data' not in st.session_state:
    st.session_state['cv_data'] = None
if 'favorites' not in st.session_state:
    st.session_state['favorites'] = []

# --- CSS Styling ---
st.markdown("""
<style>
    .stApp { background-color: #0b0e14; color: #cbd5e1; font-family: 'Segoe UI', Tahoma, sans-serif; }
    [data-testid="stSidebar"] { background-color: #111522; border-right: 1px solid #1e2436; box-shadow: 2px 0 10px rgba(0,0,0,0.2); }
    div[role="radiogroup"] > label > div:first-child { display: none !important; }
    div[role="radiogroup"] > label { padding: 14px 20px; margin-bottom: 8px; border-radius: 8px; cursor: pointer; transition: 0.3s; font-weight: 500; }
    div[role="radiogroup"] > label:hover { background-color: #1a2035; border: 1px solid #2d3748; }
    div[role="radiogroup"] > label[data-checked="true"] { background: rgba(59, 130, 246, 0.1) !important; color: #60a5fa !important; border-left: 3px solid #3b82f6; }
    
    .stButton > button { background: linear-gradient(135deg, #4f46e5, #0ea5e9); color: white; border: none; border-radius: 8px; padding: 12px 0; font-weight: bold; letter-spacing: 1px; transition: all 0.3s ease;}
    .stButton > button:hover { opacity: 0.9; transform: translateY(-2px); box-shadow: 0 4px 15px rgba(79, 70, 229, 0.4); }
    
    .glass-container { background: rgba(26, 32, 53, 0.6); backdrop-filter: blur(10px); border: 1px solid rgba(255, 255, 255, 0.05); border-radius: 16px; padding: 25px; margin-bottom: 20px; transition: transform 0.3s ease; }
    .glass-container:hover { transform: translateY(-5px); border-color: rgba(14, 165, 233, 0.3); }
    
    .metric-card { background: linear-gradient(180deg, #161c2d 0%, #111522 100%); border: 1px solid #2d3748; border-radius: 12px; padding: 20px; text-align: center; box-shadow: 0 4px 6px rgba(0,0,0,0.1); }
    .metric-title { font-size: 13px; color: #94a3b8; text-transform: uppercase; letter-spacing: 1.5px; }
    .metric-value { font-size: 28px; font-weight: 800; color: #f8fafc; margin-top: 10px; background: -webkit-linear-gradient(#fff, #94a3b8); -webkit-background-clip: text; -webkit-text-fill-color: transparent; }
    
    .job-card { background-color: #121721; border: 1px solid #1e293b; border-radius: 12px; padding: 20px; margin-bottom: 20px; }
    .job-title-row { display: flex; justify-content: space-between; align-items: center; margin-bottom: 5px; }
    .job-title { margin: 0; color: #f8fafc; font-size: 20px; font-weight: bold; }
    .match-badge { background-color: #f59e0b; color: #fff; padding: 4px 12px; border-radius: 20px; font-weight: bold; font-size: 14px; }
    .job-meta { color: #94a3b8; font-size: 13px; margin-bottom: 15px; }
    
    .new-badge { background-color: rgba(16, 185, 129, 0.15); color: #10b981; border: 1px solid #10b981; font-size: 11px; padding: 2px 8px; border-radius: 12px; margin-left: 10px; font-weight: bold;}
    .api-badge { background-color: rgba(16, 185, 129, 0.1); color: #34d399; font-size: 12px; padding: 4px 10px; border-radius: 15px; border: 1px solid #34d399; display: inline-block; margin-bottom: 15px;}
    
    .progress-container { display: flex; align-items: center; justify-content: space-between; font-size: 12px; color: #64748b; margin-bottom: 15px; border-bottom: 1px solid #1e293b; padding-bottom: 15px; }
    .progress-item { display: flex; flex-direction: column; width: 30%; }
    .progress-bar-bg { width: 100%; height: 4px; background-color: #1e293b; border-radius: 4px; margin-top: 5px; overflow: hidden; }
    .progress-fill { height: 100%; border-radius: 4px; }
    
    .why-matched { background-color: rgba(59, 130, 246, 0.05); border-left: 3px solid #3b82f6; padding: 12px 15px; border-radius: 0 8px 8px 0; margin-bottom: 15px; font-size: 13px; color: #cbd5e1; }
    
    .skill-pill-green { background: rgba(16, 185, 129, 0.1); border: 1px solid #10b981; color: #10b981; padding: 4px 12px; border-radius: 20px; font-size: 12px; margin: 0 6px 8px 0; display: inline-block; }
    .skill-pill-red { background: rgba(239, 68, 68, 0.1); border: 1px solid #ef4444; color: #ef4444; padding: 4px 12px; border-radius: 20px; font-size: 12px; margin: 0 6px 8px 0; display: inline-block; }
    .skill-pill-orange { background: rgba(245, 158, 11, 0.1); border: 1px solid #f59e0b; color: #f59e0b; padding: 4px 12px; border-radius: 20px; font-size: 12px; margin: 0 6px 8px 0; display: inline-block; }
    
    .cv-paper { background-color: #ffffff; color: #1a202c; padding: 40px; border-radius: 4px; box-shadow: 0 10px 25px rgba(0,0,0,0.5); min-height: 800px; font-family: 'Arial', sans-serif; }
</style>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("<h2 style='color: white; margin-bottom: 0; font-weight: 800;'>JobMatch <span style='color: #3b82f6;'>AI</span></h2>", unsafe_allow_html=True)
    st.caption("NEXT-GEN CAREER PLATFORM")
    st.markdown("<br>", unsafe_allow_html=True)
    menu = st.radio("", ["Home", "AI Recommendations", "CV Analyzer", "Live Jobs", "Saved Jobs", "CV Builder", "Domain AI"], label_visibility="collapsed")
    st.markdown("<br><br><br><br>", unsafe_allow_html=True)
    st.markdown("---")
    st.markdown("<p style='color: #64748b; font-size: 12px;'>API Access <br> <span style='color: #10b981;'>● v5.0.0 Live</span></p>", unsafe_allow_html=True)

# --- Helper Functions ---
def is_job_new(date_str):
    if not date_str or pd.isna(date_str):
        return False
    d_str = str(date_str).lower()
    
    if "hour" in d_str or "minute" in d_str or "second" in d_str or "today" in d_str or "just now" in d_str:
        return True
    if "day" in d_str:
        nums = re.findall(r'\d+', d_str)
        if nums and int(nums[0]) <= 7:
            return True
            
    try:
        parsed_date = pd.to_datetime(date_str)
        if (pd.Timestamp.now() - parsed_date).days <= 7:
            return True
    except:
        pass
    return False

def extract_contact_info(text):
    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
    phone_match = re.search(r'(\+?\d{1,3}[\s-]?)?\(?\d{2,4}\)?[\s-]?\d{3,4}[\s-]?\d{3,4}', text)
    return {
        "email": email_match.group(0) if email_match else "Not Found",
        "phone": phone_match.group(0) if phone_match else "Not Found"
    }

def toggle_favorite(job):
    job_key = f"{job.get('Title', '')}_{job.get('Company', '')}"
    is_saved = any(f.get('Title') == job.get('Title') and f.get('Company') == job.get('Company') for f in st.session_state['favorites'])
    if is_saved:
        st.session_state['favorites'] = [f for f in st.session_state['favorites'] if not (f.get('Title') == job.get('Title') and f.get('Company') == job.get('Company'))]
    else:
        st.session_state['favorites'].append(job)

def render_job_card(job, cv_role="Software Engineer", skill_pct=0, exp_pct=0, level_pct=0, show_match_bars=True):
    match_score = job.get('Match_Score', 'N/A')
    is_new = is_job_new(job.get('Date', '')) or is_job_new(job.get('Posted', ''))
    new_badge = "<span class='new-badge'>NEW</span>" if is_new else ""
    source = job.get('Source', 'Live API')
    
    match_badge_html = ""
    if show_match_bars and match_score != 'N/A':
        match_badge_html = f"<span class='match-badge'>⭐ {match_score}{'%' if isinstance(match_score, int) else ''} Match</span>"
    
    # تمت إزالة جميع المسافات البادئة لتجنب تفسيرها كـ Markdown Code
    html = f"""
<div class="job-card">
<div class="job-title-row">
<h3 class="job-title">{job.get('Title', 'Unknown Title')} {new_badge}</h3>
{match_badge_html}
</div>
<div class="job-meta">
🏢 {job.get('Company', 'Unknown')} &nbsp;|&nbsp; 📍 {job.get('Location', 'Unknown')}
</div>
<span class="api-badge">● {source}</span>
"""
    
    if show_match_bars:
        html += f"""
<div class="progress-container">
<div class="progress-item">
<span style="display:flex; justify-content:space-between;"><span>Skills ({int(skill_pct)}%)</span> </span>
<div class="progress-bar-bg"><div class="progress-fill" style="width: {skill_pct}%; background-color: #3b82f6;"></div></div>
</div>
<div class="progress-item">
<span style="display:flex; justify-content:space-between;"><span>Experience ({int(exp_pct)}%)</span> </span>
<div class="progress-bar-bg"><div class="progress-fill" style="width: {exp_pct}%; background-color: #10b981;"></div></div>
</div>
<div class="progress-item">
<span style="display:flex; justify-content:space-between;"><span>Domain Match ({int(level_pct)}%)</span> </span>
<div class="progress-bar-bg"><div class="progress-fill" style="width: {level_pct}%; background-color: #f59e0b;"></div></div>
</div>
</div>
<div class="why-matched">
<strong>🪄 Why matched:</strong> Matches your skills in {cv_role}. Relevant to your career direction. 
{'Opportunity to learn new technologies.' if job.get('Missing_Skills') else 'Excellent technical alignment.'}
</div>
"""
    
    desc = str(job.get('Description', ''))
    short_desc = desc[:250] + "..." if len(desc) > 250 else desc
    if short_desc:
        html += f"<p style='color: #94a3b8; font-size: 13px; line-height: 1.5; margin-bottom: 15px;'>{short_desc}</p>"

    if show_match_bars:
        matched_skills = job.get('Matched_Skills', [])
        missing_skills = job.get('Missing_Skills', [])
        
        skills_html = "<div style='margin-bottom: 15px;'>"
        for s in matched_skills[:10]:
            skills_html += f"<span class='skill-pill-green'>✓ {s.title()}</span>"
        for s in missing_skills[:8]:
            skills_html += f"<span class='skill-pill-red'>× {s.title()}</span>"
        skills_html += "</div>"
        
        if missing_skills:
            skills_html += f"<div><span style='color: #f59e0b; font-size: 13px; margin-right: 10px;'>🎓 Recommended Skills to Learn:</span><br>"
            for s in missing_skills[:5]:
                skills_html += f"<span class='skill-pill-orange'>{s.title()}</span>"
            skills_html += "</div>"
            
        html += skills_html

    html += "</div>"
    st.markdown(html, unsafe_allow_html=True)
    
    col1, col2, col3 = st.columns([2, 1, 7])
    with col1:
        if 'Apply_Link' in job and str(job['Apply_Link']).startswith('http'):
            st.link_button("↗ Apply Now", job['Apply_Link'], use_container_width=True)
        else:
            st.info("No Link")
    with col2:
        is_saved = any(f.get('Title') == job.get('Title') and f.get('Company') == job.get('Company') for f in st.session_state['favorites'])
        if st.button("Unsave" if is_saved else "Save", key=f"fav_{job.get('Title')}_{job.get('Company')}", use_container_width=True):
            toggle_favorite(job)
            st.rerun()
    st.write("<br>", unsafe_allow_html=True)

# ======================== الصفحات ======================== #

if menu == "Home":
    st.markdown("<h1 style='font-weight: 800; letter-spacing: -1px;'>Welcome to JobMatch AI</h1>", unsafe_allow_html=True)
    st.markdown("<p style='font-size: 18px; color: #94a3b8; max-width: 800px; line-height: 1.6;'>Your intelligent career copilot. We bridge the gap between your skills and market demands using advanced Machine Learning.</p>", unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    
    col1, col2, col3 = st.columns(3)
    with col1: st.markdown("""<div class="metric-card"><div class="metric-title">Jobs Analyzed</div><div class="metric-value">353,052+</div></div>""", unsafe_allow_html=True)
    with col2: st.markdown("""<div class="metric-card"><div class="metric-title">Active Domains</div><div class="metric-value">19 Domains</div></div>""", unsafe_allow_html=True)
    with col3: st.markdown(f"""<div class="metric-card"><div class="metric-title">Saved Jobs</div><div class="metric-value">{len(st.session_state['favorites'])}</div></div>""", unsafe_allow_html=True)
    
    st.markdown("<br><br><h3 style='color: #e2e8f0;'>Quick Actions</h3>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        st.info("**Analyze my CV:** Upload your resume to extract pure technical skills and find matching jobs.")
        st.info("**AI Match:** Manually enter your tech stack and see matching jobs instantly.")
    with c2:
        st.success("**Build CV:** Use our ATS-friendly builder to create a winning resume.")
        st.success("**Explore Domains:** See what skills are trending in the tech industry.")

elif menu == "AI Recommendations" or menu == "AI Match (Manual)":
    st.header("AI Job Recommendations")
    st.caption("/ Real-Time Job Market Matching based on your stack")
    
    with st.container(border=True):
        manual_skills = st.text_area("Your Technical Stack", placeholder="e.g., python, pandas, scikit-learn, aws...")
        col1, col2, col3 = st.columns(3)
        with col1: exp_years = st.number_input("Experience (years)", value=3.0, step=0.5)
        with col2: career_level = st.selectbox("Career Level", ["Entry Level", "Experienced", "Manager"])
        with col3: country_code = st.text_input("Country Code", value="eg")
            
        find_matches = st.button("Generate Recommendations", type="primary", use_container_width=True)

    if find_matches and manual_skills:
        with st.spinner("Analyzing market..."):
            cleaned_manual = clean_text(manual_skills)
            _, _, skills_vectorizer, _ = load_models()
            extracted_list = [s.strip().lower() for s in manual_skills.split(",") if s.strip()]
            if skills_vectorizer:
                try:
                    vec = skills_vectorizer.transform([cleaned_manual])
                    extracted_list = skills_vectorizer.inverse_transform(vec)[0].tolist()
                except: pass
                
            cv_data = {
                "skills_text": " ; ".join(extracted_list) if extracted_list else cleaned_manual,
                "skills": extracted_list,
                "experience_level": exp_years,
                "career_level": career_level,
                "role": "Tech Professional"
            }
            live_jobs = fetch_live_jobs(manual_skills, country_code)
            if live_jobs:
                jobs_df = pd.DataFrame(live_jobs)
                ranked_jobs = calculate_match_score(cv_data, jobs_df)
                for job in ranked_jobs[:10]:
                    render_job_card(
                        job, 
                        cv_role="Tech Professional",
                        skill_pct=min(job.get('Skill_Similarity', 0)*100, 100), 
                        exp_pct=min(job.get('Exp_Score', 0)*100, 100), 
                        level_pct=min(job.get('Level_Score', 0)*100, 100), 
                        show_match_bars=True
                    )
            else:
                st.warning("No jobs found matching your specific skills.")

elif menu == "CV Analyzer":
    st.header("AI CV Analyzer")
    st.caption("Upload your resume for deep semantic analysis and job matching.")
    
    col_up, col_ct = st.columns([3, 1])
    with col_up: uploaded_file = st.file_uploader("Drop your CV here", type=["pdf", "docx", "txt"])
    with col_ct: search_country = st.text_input("Search Jobs in (Country)", value="eg")
    
    if uploaded_file and st.button("Analyze My CV", type="primary", use_container_width=True):
        with st.spinner("Extracting features..."):
            st.session_state['cv_data'] = analyze_cv(uploaded_file)
    
    if st.session_state['cv_data'] is not None:
        cv_data = st.session_state['cv_data']
        contact_info = extract_contact_info(cv_data['raw_text'])
        
        st.markdown(f"""
        <div class="glass-container">
        <h4 style="margin-top:0;">📇 Contact Info Extracted</h4>
        <p style="color:#94a3b8; margin:0;">📧 {contact_info['email']} <br> 📱 {contact_info['phone']}</p>
        </div>
        """, unsafe_allow_html=True)
        
        col1, col2, col3, col4 = st.columns(4)
        with col1: st.markdown(f"""<div class="metric-card"><div class="metric-title">Skills Found</div><div class="metric-value">{len(cv_data['skills'])}</div></div>""", unsafe_allow_html=True)
        with col2: st.markdown(f"""<div class="metric-card"><div class="metric-title">Est. Experience</div><div class="metric-value">{cv_data['experience_level']} yrs</div></div>""", unsafe_allow_html=True)
        with col3: st.markdown(f"""<div class="metric-card"><div class="metric-title">Career Level</div><div class="metric-value" style="font-size: 20px;">{cv_data['career_level']}</div></div>""", unsafe_allow_html=True)
        with col4: st.markdown(f"""<div class="metric-card"><div class="metric-title">Top Domain</div><div class="metric-value" style="font-size: 16px;">{cv_data['role']}</div></div>""", unsafe_allow_html=True)
        
        st.markdown("<br><h4 style='color: #e2e8f0;'>Extracted Technical Skills</h4>", unsafe_allow_html=True)
        if cv_data['skills']:
            skills_html = "".join([f"<span class='skill-pill-green'>{skill.title()}</span>" for skill in cv_data['skills']])
        else:
            skills_html = "<p style='color:#94a3b8;'>No standard technical skills detected.</p>"
        st.markdown(f"<div class='glass-container'>{skills_html}</div>", unsafe_allow_html=True)

        st.markdown("<hr>", unsafe_allow_html=True)
        
        if st.button("🔍 Find Jobs Matching My CV", type="primary", use_container_width=True):
            with st.spinner("Fetching and matching the best jobs for you..."):
                search_query = cv_data['role']
                live_jobs = fetch_live_jobs(search_query, search_country)
                if live_jobs:
                    jobs_df = pd.DataFrame(live_jobs)
                    ranked_jobs = calculate_match_score(cv_data, jobs_df)
                    st.success(f"Found {len(ranked_jobs)} matched jobs!")
                    for job in ranked_jobs[:10]:
                        skill_pct = min(job.get('Skill_Similarity', 0)*100, 100)
                        exp_pct = min(job.get('Exp_Score', 0)*100, 100)
                        level_pct = min(job.get('Level_Score', 0)*100, 100)
                        render_job_card(job, cv_role=cv_data['role'], skill_pct=skill_pct, exp_pct=exp_pct, level_pct=level_pct, show_match_bars=True)
                else:
                    st.warning("No jobs found matching your specific profile right now.")

elif menu == "Live Jobs":
    st.header("Live Job Search")
    st.caption("Browse all latest jobs directly.")
    search_q = st.text_input("Job Title", value="Software Engineer")
    if st.button("Search"):
        jobs = fetch_live_jobs(search_q, "eg")
        for job in jobs[:15]:
            # استخدام التصميم الشيك الأساسي لكن بدون تفاصيل الماتش (show_match_bars=False)
            render_job_card(job, show_match_bars=False)

elif menu == "Saved Jobs":
    st.header("Saved Jobs")
    if not st.session_state['favorites']:
        st.info("No saved jobs yet.")
    else:
        for job in st.session_state['favorites']:
            render_job_card(job, show_match_bars=False)

elif menu == "CV Builder":
    st.header("ATS-Optimized CV Builder")
    st.caption("Build a comprehensive, ATS-friendly resume.")
    
    def_name = ""
    def_email = ""
    def_phone = ""
    def_title = ""
    def_skills = ""
    
    if st.session_state.get('cv_data'):
        c_data = st.session_state['cv_data']
        def_title = c_data.get('role', '')
        def_skills = ", ".join(c_data.get('skills', []))
        c_info = extract_contact_info(c_data.get('raw_text', ''))
        if c_info['email'] != "Not Found": def_email = c_info['email']
        if c_info['phone'] != "Not Found": def_phone = c_info['phone']
        lines = c_data.get('raw_text', '').split('\n')
        if lines:
            def_name = lines[0].strip()[:30]

    def generate_docx_cv(name, title, email, phone, linkedin, summary, experience, projects, education, skills, languages):
        doc = docx.Document()
        
        name_para = doc.add_paragraph()
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_run = name_para.add_run(name)
        name_run.bold = True
        name_run.font.size = Pt(22)
        
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_para.add_run(f"{email} | {phone} | {linkedin}").font.size = Pt(10)
        
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_para.add_run(title).font.size = Pt(12)

        def add_section(heading, text_content):
            if text_content.strip():
                doc.add_heading(heading, level=2)
                for line in text_content.split('\n'):
                    if line.strip():
                        if heading in ['EXPERIENCE', 'PROJECTS', 'EDUCATION']:
                            doc.add_paragraph(line.strip(), style='List Bullet')
                        else:
                            doc.add_paragraph(line.strip())

        add_section('PROFESSIONAL SUMMARY', summary)
        add_section('EXPERIENCE', experience)
        add_section('PROJECTS', projects)
        add_section('EDUCATION', education)
        add_section('TECHNICAL SKILLS', skills)
        add_section('LANGUAGES', languages)
        
        bio = io.BytesIO()
        doc.save(bio)
        return bio.getvalue()

    col_form, col_preview = st.columns([1, 1.2])
    
    with col_form:
        cv_name = st.text_input("Full Name", value=def_name, placeholder="e.g., Omar Tarek")
        cv_title = st.text_input("Professional Title", value=def_title, placeholder="e.g., AI Engineer")
        col_c1, col_c2 = st.columns(2)
        with col_c1: cv_email = st.text_input("Email", value=def_email)
        with col_c2: cv_phone = st.text_input("Phone", value=def_phone)
        cv_linkedin = st.text_input("LinkedIn / Portfolio", "linkedin.com/in/omar")
        
        cv_summary = st.text_area("Professional Summary", "Passionate AI Engineer...", height=80)
        cv_exp = st.text_area("Work Experience", "Company Name - Role (Date)\n- Achievement 1\n- Achievement 2", height=100)
        cv_proj = st.text_area("Projects", "Project Name\n- Description and technologies used (Python, NLP).", height=100)
        cv_edu = st.text_area("Education", "B.Sc. Artificial Intelligence Engineering\nTanta University | 2026", height=80)
        cv_skills = st.text_area("Technical Skills", value=def_skills if def_skills else "Python, Machine Learning, Scikit-learn, SQL", height=80)
        cv_lang = st.text_area("Languages", "Arabic (Native)\nEnglish (Fluent)", height=60)
            
    with col_preview:
        html_preview = f"""
<div class="cv-paper">
<h1 style='text-align: center; color: #000; margin-bottom: 0px;'>{cv_name if cv_name else 'Your Name'}</h1>
<p style='text-align: center; color: #333; font-size: 12px; margin-top: 5px;'>{cv_email if cv_email else 'Email'} | {cv_phone if cv_phone else 'Phone'} | {cv_linkedin}</p>
<p style='text-align: center; color: #000; font-size: 14px; font-weight: bold;'>{cv_title if cv_title else 'Professional Title'}</p>

<h3 style='border-bottom: 1px solid #000; font-size: 14px; margin-top: 20px;'>SUMMARY</h3>
<p style='font-size: 13px; line-height: 1.5;'>{cv_summary}</p>

<h3 style='border-bottom: 1px solid #000; font-size: 14px; margin-top: 15px;'>EXPERIENCE</h3>
<ul style='font-size: 13px; line-height: 1.5;'>{"".join([f"<li>{line}</li>" for line in cv_exp.split(chr(10)) if line.strip()])}</ul>

<h3 style='border-bottom: 1px solid #000; font-size: 14px; margin-top: 15px;'>PROJECTS</h3>
<ul style='font-size: 13px; line-height: 1.5;'>{"".join([f"<li>{line}</li>" for line in cv_proj.split(chr(10)) if line.strip()])}</ul>

<h3 style='border-bottom: 1px solid #000; font-size: 14px; margin-top: 15px;'>EDUCATION</h3>
<ul style='font-size: 13px; line-height: 1.5;'>{"".join([f"<li>{line}</li>" for line in cv_edu.split(chr(10)) if line.strip()])}</ul>

<h3 style='border-bottom: 1px solid #000; font-size: 14px; margin-top: 15px;'>SKILLS</h3>
<p style='font-size: 13px; line-height: 1.5;'>{cv_skills}</p>

<h3 style='border-bottom: 1px solid #000; font-size: 14px; margin-top: 15px;'>LANGUAGES</h3>
<p style='font-size: 13px; line-height: 1.5;'>{cv_lang.replace(chr(10), '<br>')}</p>
</div>
"""
        st.markdown(html_preview, unsafe_allow_html=True)
        
        st.write("")
        docx_data = generate_docx_cv(cv_name, cv_title, cv_email, cv_phone, cv_linkedin, cv_summary, cv_exp, cv_proj, cv_edu, cv_skills, cv_lang)
        st.download_button(
            label="📥 Download ATS CV (.docx)",
            data=docx_data,
            file_name=f"{cv_name.replace(' ', '_')}_CV.docx" if cv_name else "CV.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )

elif menu == "Domain AI":
    st.header("Domain Intelligence")
    st.info("Select a domain to see market insights.")
    domain_sel = st.selectbox("Select a Domain", ["Machine Learning / AI", "Data Science & Analytics", "Data Engineering", "Software Engineering", "Cloud & DevOps", "Cybersecurity"])
    
    st.markdown(f"<br><h3 style='color: #f8fafc;'>Market Insights: {domain_sel}</h3>", unsafe_allow_html=True)
    if domain_sel == "Machine Learning / AI":
        st.info("Focuses on building intelligent systems, deep learning models, NLP, computer vision, and generative AI applications.")
        st.markdown("<span class='skill-pill-green'>Python</span> <span class='skill-pill-green'>PyTorch</span> <span class='skill-pill-green'>TensorFlow</span> <span class='skill-pill-green'>NLP</span> <span class='skill-pill-green'>Docker</span>", unsafe_allow_html=True)
    elif domain_sel == "Data Science & Analytics":
        st.info("Focuses on extracting actionable insights from massive datasets using statistical modeling and BI tools.")
        st.markdown("<span class='skill-pill-green'>Python</span> <span class='skill-pill-green'>SQL</span> <span class='skill-pill-green'>Power BI</span> <span class='skill-pill-green'>Tableau</span> <span class='skill-pill-green'>Pandas</span>", unsafe_allow_html=True)
    elif domain_sel == "Data Engineering":
        st.info("Focuses on designing and maintaining scalable data pipelines, ETL workflows, and cloud data warehouses.")
        st.markdown("<span class='skill-pill-green'>SQL</span> <span class='skill-pill-green'>Spark</span> <span class='skill-pill-green'>Airflow</span> <span class='skill-pill-green'>Snowflake</span> <span class='skill-pill-green'>Kafka</span>", unsafe_allow_html=True)
    elif domain_sel == "Software Engineering":
        st.info("Core development of scalable web, mobile, and backend systems, microservices, and RESTful APIs.")
        st.markdown("<span class='skill-pill-green'>JavaScript</span> <span class='skill-pill-green'>React</span> <span class='skill-pill-green'>Node.js</span> <span class='skill-pill-green'>Python</span> <span class='skill-pill-green'>Go</span> <span class='skill-pill-green'>C++</span>", unsafe_allow_html=True)
    elif domain_sel == "Cloud & DevOps":
        st.info("Automating infrastructure, CI/CD deployment pipelines, container orchestration, and cloud architecture.")
        st.markdown("<span class='skill-pill-green'>AWS</span> <span class='skill-pill-green'>Docker</span> <span class='skill-pill-green'>Kubernetes</span> <span class='skill-pill-green'>Terraform</span> <span class='skill-pill-green'>CI/CD</span>", unsafe_allow_html=True)
    else:
        st.info("Protecting networks, systems, and applications from digital attacks and vulnerabilities.")
        st.markdown("<span class='skill-pill-green'>Network Security</span> <span class='skill-pill-green'>Linux</span> <span class='skill-pill-green'>Penetration Testing</span> <span class='skill-pill-green'>SIEM</span>", unsafe_allow_html=True)